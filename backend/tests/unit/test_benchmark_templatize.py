"""Unit tests for the assisted templating benchmark script."""

import io
import json
from pathlib import Path

from docx import Document
from pydantic_ai.models.test import TestModel

from scripts.benchmark_templatize import (
    aggregate,
    benchmark_templates,
    docx_text,
    mapping_agreement,
    text_similarity,
)


def _docx(path: Path, *paragraphs: str) -> Path:
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    document.save(str(path))
    return path


def test_mapping_agreement_scores_pairs() -> None:
    ref = [
        {"find": "Jean", "placeholder": "first_name"},
        {"find": "Dupont", "placeholder": "last_name"},
    ]
    cand = [
        {"find": "Jean", "placeholder": "first_name"},
        {"find": "Paris", "placeholder": "location"},
    ]
    assert mapping_agreement(cand, ref) == {"mapping_precision": 0.5, "mapping_recall": 0.5}


def test_mapping_agreement_both_empty_is_perfect() -> None:
    assert mapping_agreement([], []) == {"mapping_precision": 1.0, "mapping_recall": 1.0}


def test_text_similarity() -> None:
    assert text_similarity("a\nb", "a\nb") == 1.0
    assert text_similarity("a\nb", "c\nd") == 0.0


def test_docx_text_reads_paragraphs_and_tables() -> None:
    document = Document()
    document.add_paragraph("Titre")
    document.add_table(rows=1, cols=1).rows[0].cells[0].text = "Cellule"
    buffer = io.BytesIO()
    document.save(buffer)
    assert docx_text(buffer.getvalue()) == "Titre\nCellule"


async def test_benchmark_writes_outputs_and_scores_against_reference(tmp_path: Path) -> None:
    templates = tmp_path / "templates"
    templates.mkdir()
    template = _docx(templates / "cabinet.docx", "Dossier de competences")
    model = TestModel(custom_output_args={"operations": []})

    reference = await benchmark_templates(model, [template], tmp_path / "ref", None)
    assert reference[0]["status"] == "ok"
    assert reference[0]["llm_requests"] == 1
    assert reference[0]["cost_usd"] is None  # test model has no price
    assert (tmp_path / "ref" / "cabinet.render.txt").read_text() == "Dossier de competences"
    report = json.loads((tmp_path / "ref" / "cabinet.report.json").read_text())
    assert report["metrics"]["status"] == "ok"

    run = await benchmark_templates(model, [template], tmp_path / "run", tmp_path / "ref")
    assert run[0]["render_similarity"] == 1.0
    assert run[0]["mapping_recall"] == 1.0


def test_aggregate_cost_per_success() -> None:
    results = [
        {"status": "ok", "llm_requests": 1, "cost_usd": 0.4},
        {"status": "render_error", "llm_requests": 2, "cost_usd": 0.6},
    ]
    summary = aggregate(results)
    assert summary["success_rate"] == 0.5
    assert summary["first_pass_success"] == 1
    assert summary["cost_per_success_usd"] == 1.0

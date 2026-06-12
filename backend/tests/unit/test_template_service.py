# backend/tests/unit/test_template_service.py
"""Unit tests for render-based template validation."""

from pathlib import Path

from docx import Document

from services.documents.template_service import TemplateValidation, validate_template


def _docx(tmp_path: Path, paragraphs: list[str]) -> str:
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    path = tmp_path / "t.docx"
    doc.save(str(path))
    return str(path)


def test_template_with_known_placeholders_is_valid(tmp_path: Path) -> None:
    path = _docx(tmp_path, ["{{first_name}} {{last_name}}", "{{availability_label}}"])
    result = validate_template(path, ["{{first_name}}", "{{last_name}}", "{{availability_label}}"])
    assert result == TemplateValidation(True, [], None)


def test_template_with_spaces_and_doc_loops_is_valid(tmp_path: Path) -> None:
    # Conforme a docs/template-syntax.md : espaces toleres, boucles edu/cert/lang.
    path = _docx(
        tmp_path,
        [
            "{{ first_name }}",
            "{%p for edu in educations %}",
            "{{edu.degree}}",
            "{%p endfor %}",
        ],
    )
    result = validate_template(path, ["{{ first_name }}"])
    assert result.is_valid is True
    assert result.unknown_placeholders == []


def test_unknown_placeholder_is_warning_not_blocking(tmp_path: Path) -> None:
    path = _docx(tmp_path, ["{{NOM}} {{first_name}}"])
    result = validate_template(path, ["{{NOM}}", "{{first_name}}"])
    assert result.is_valid is True
    assert result.unknown_placeholders == ["{{NOM}}"]
    assert result.validation_error is None


def test_broken_jinja_is_invalid(tmp_path: Path) -> None:
    path = _docx(tmp_path, ["{%p for exp in experiences %}", "{{exp.role}}"])  # endfor manquant
    result = validate_template(path, [])
    assert result.is_valid is False
    assert result.validation_error is not None


def test_unreadable_file_is_invalid(tmp_path: Path) -> None:
    path = tmp_path / "not_a_docx.docx"
    path.write_bytes(b"garbage")
    result = validate_template(str(path), [])
    assert result.is_valid is False
    assert result.validation_error is not None

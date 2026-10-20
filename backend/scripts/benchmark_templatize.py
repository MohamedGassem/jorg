#!/usr/bin/env python3
"""Benchmark assisted templating on a folder of company Word templates.

One run = one (model, effort) pair. Each template goes through the production
pipeline (run_templatize_pipeline). The script records render success, LLM
requests, tokens and cost, and writes for each template:
- <name>.docx         the templatized document
- <name>.report.json  the pipeline report plus the run metrics
- <name>.render.txt   the text of the mock render (only when it renders)

Once a run has been checked by hand, its output folder becomes the reference:
later runs passed --reference are scored against it (mock render similarity,
mapping precision/recall).

Usage (from backend/):
    LLM_API_KEY=... uv run python -m scripts.benchmark_templatize \\
        --templates ~/templates_clients --out benchmarks/opus55-high

    LLM_API_KEY=... uv run python -m scripts.benchmark_templatize \\
        --templates ~/templates_clients --out benchmarks/sonnet55-medium \\
        --model anthropic:claude-sonnet-5-5 --effort medium \\
        --reference benchmarks/opus55-high
"""

from __future__ import annotations

import argparse
import asyncio
import difflib
import io
import json
import os
import sys
import time
from pathlib import Path
from typing import Any

from docx import Document
from pydantic_ai.messages import ModelMessage, ModelResponse
from pydantic_ai.models import Model, ModelRequestParameters
from pydantic_ai.models.wrapper import WrapperModel
from pydantic_ai.settings import ModelSettings

from core.config import get_settings
from services.documents.builtin_template_service import render_mock_preview_from_path
from services.documents.templatize_service import run_templatize_pipeline
from services.llm.client import get_llm_model
from services.llm.templatize import TemplatizeLLMError


class UsageRecorder(WrapperModel):
    """Keeps every LLM response so the run can report tokens and cost."""

    def __init__(self, wrapped: Model) -> None:
        super().__init__(wrapped)
        self.responses: list[ModelResponse] = []

    async def request(
        self,
        messages: list[ModelMessage],
        model_settings: ModelSettings | None,
        model_request_parameters: ModelRequestParameters,
    ) -> ModelResponse:
        response = await super().request(messages, model_settings, model_request_parameters)
        self.responses.append(response)
        return response


def usage_summary(responses: list[ModelResponse]) -> dict[str, Any]:
    """Token totals and cost in USD (None when the model has no known price)."""
    cost: float | None = 0.0
    for response in responses:
        try:
            cost = None if cost is None else cost + float(response.cost().total_price)
        except LookupError:
            cost = None
    return {
        "llm_requests": len(responses),
        "input_tokens": sum(r.usage.input_tokens for r in responses),
        "output_tokens": sum(r.usage.output_tokens for r in responses),
        "cost_usd": cost,
    }


def docx_text(docx_bytes: bytes) -> str:
    """Plain text of a docx: body paragraphs, then table cells, one line each."""
    document = Document(io.BytesIO(docx_bytes))
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.extend(cell.text for cell in row.cells)
    return "\n".join(line for line in lines if line.strip())


def text_similarity(candidate: str, reference: str) -> float:
    """Line-level similarity ratio between two mock renders (1.0 = identical)."""
    return difflib.SequenceMatcher(None, candidate.splitlines(), reference.splitlines()).ratio()


def mapping_agreement(
    candidate: list[dict[str, str]], reference: list[dict[str, str]]
) -> dict[str, float]:
    """Precision/recall of (find, placeholder) pairs against the reference run."""
    cand = {(m["find"], m["placeholder"]) for m in candidate}
    ref = {(m["find"], m["placeholder"]) for m in reference}
    common = len(cand & ref)
    return {
        "mapping_precision": common / len(cand) if cand else float(not ref),
        "mapping_recall": common / len(ref) if ref else float(not cand),
    }


def compare_to_reference(
    name: str, report: dict[str, Any], render_text: str | None, reference_dir: Path
) -> dict[str, Any]:
    ref_report_path = reference_dir / f"{name}.report.json"
    if not ref_report_path.exists():
        return {"reference": "missing"}
    ref_report = json.loads(ref_report_path.read_text(encoding="utf-8"))
    scores: dict[str, Any] = mapping_agreement(report["mappings"], ref_report["mappings"])
    ref_render_path = reference_dir / f"{name}.render.txt"
    if render_text is not None and ref_render_path.exists():
        ref_render = ref_render_path.read_text(encoding="utf-8")
        scores["render_similarity"] = text_similarity(render_text, ref_render)
    else:
        scores["render_similarity"] = 0.0
    return scores


async def benchmark_template(
    model: Model, template: Path, out_dir: Path, reference_dir: Path | None
) -> dict[str, Any]:
    recorder = UsageRecorder(model)
    name = template.stem
    started = time.perf_counter()
    result: dict[str, Any] = {"template": template.name}
    report: dict[str, Any] = {"mappings": [], "warnings": [], "rejected": []}
    render_text: str | None = None
    try:
        outcome = await run_templatize_pipeline(recorder, str(template))
    except TemplatizeLLMError as exc:
        result.update(status="llm_error", error=str(exc))
    except Exception as exc:  # recorded per template so one bad file does not stop the run
        result.update(status="error", error=f"{exc.__class__.__name__}: {exc}")
    else:
        report = outcome.report
        docx_path = out_dir / f"{name}.docx"
        docx_path.write_bytes(outcome.docx_bytes)
        if outcome.render_error is None:
            render_text = docx_text(render_mock_preview_from_path(str(docx_path)))
            (out_dir / f"{name}.render.txt").write_text(render_text, encoding="utf-8")
            result["status"] = "ok"
        else:
            result.update(status="render_error", error=outcome.render_error)

    result.update(usage_summary(recorder.responses))
    result["duration_s"] = round(time.perf_counter() - started, 1)
    result["mappings"] = len(report["mappings"])
    result["warnings"] = len(report["warnings"])
    result["rejected"] = len(report["rejected"])
    if reference_dir is not None:
        result.update(compare_to_reference(name, report, render_text, reference_dir))

    (out_dir / f"{name}.report.json").write_text(
        json.dumps({**report, "metrics": result}, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return result


async def benchmark_templates(
    model: Model, templates: list[Path], out_dir: Path, reference_dir: Path | None
) -> list[dict[str, Any]]:
    """Run the pipeline on each template sequentially (keeps rate limits and logs simple)."""
    out_dir.mkdir(parents=True, exist_ok=True)
    results = []
    for template in templates:
        print(f"-> {template.name}", flush=True)
        results.append(await benchmark_template(model, template, out_dir, reference_dir))
    return results


def aggregate(results: list[dict[str, Any]]) -> dict[str, Any]:
    ok = [r for r in results if r["status"] == "ok"]
    costs = [r["cost_usd"] for r in results]
    total_cost = None if any(c is None for c in costs) else sum(costs)
    summary: dict[str, Any] = {
        "templates": len(results),
        "success": len(ok),
        "success_rate": len(ok) / len(results) if results else 0.0,
        "first_pass_success": sum(1 for r in ok if r["llm_requests"] == 1),
        "total_cost_usd": total_cost,
        "cost_per_success_usd": total_cost / len(ok) if ok and total_cost is not None else None,
    }
    scored = [r for r in results if "render_similarity" in r]
    if scored:
        summary["mean_render_similarity"] = sum(r["render_similarity"] for r in scored) / len(
            scored
        )
        summary["mean_mapping_recall"] = sum(r["mapping_recall"] for r in scored) / len(scored)
    return summary


def print_table(results: list[dict[str, Any]], summary: dict[str, Any]) -> None:
    print(f"\n{'template':40} {'status':13} {'req':>3} {'cost $':>7} {'sim':>5}")
    for r in results:
        cost = "-" if r["cost_usd"] is None else f"{r['cost_usd']:.3f}"
        sim = f"{r['render_similarity']:.2f}" if "render_similarity" in r else "-"
        print(f"{r['template'][:40]:40} {r['status']:13} {r['llm_requests']:>3} {cost:>7} {sim:>5}")
    print("\n" + json.dumps(summary, indent=2))


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    parser.add_argument("--templates", type=Path, required=True, help="folder of .docx templates")
    parser.add_argument("--out", type=Path, required=True, help="output folder for this run")
    parser.add_argument("--model", help="Pydantic AI model, e.g. anthropic:claude-sonnet-5-5")
    parser.add_argument("--effort", choices=["minimal", "low", "medium", "high", "xhigh"])
    parser.add_argument("--reference", type=Path, help="output folder of a validated run")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    if not args.templates.is_dir():
        print(f"Templates folder not found: {args.templates}", file=sys.stderr)
        return 2
    templates = sorted(args.templates.glob("*.docx"))
    if not templates:
        print(f"No .docx file in {args.templates}", file=sys.stderr)
        return 2
    if args.reference is not None and not args.reference.is_dir():
        print(f"Reference folder not found: {args.reference}", file=sys.stderr)
        return 2

    # CLI flags override the env so one run = one (model, effort) pair.
    if args.model:
        os.environ["LLM_MODEL"] = args.model
    if args.effort:
        os.environ["LLM_EFFORT"] = args.effort
    get_settings.cache_clear()
    settings = get_settings()
    model = get_llm_model()
    if model is None:
        print("LLM_API_KEY (or ANTHROPIC_API_KEY) is not set", file=sys.stderr)
        return 2

    results = asyncio.run(benchmark_templates(model, templates, args.out, args.reference))
    summary = {"model": settings.llm_model, "effort": settings.llm_effort, **aggregate(results)}
    (args.out / "summary.json").write_text(
        json.dumps({"summary": summary, "results": results}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print_table(results, summary)
    return 0


if __name__ == "__main__":
    sys.exit(main())
